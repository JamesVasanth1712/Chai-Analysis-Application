import io
import json
import re
from typing import Any, Dict, List, Tuple
import pandas as pd
import structlog

logger = structlog.get_logger()


def parse_json(contents: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        data = json.loads(contents.decode("utf-8-sig", errors="replace"))
        if isinstance(data, list):
            if len(data) > 0 and isinstance(data[0], dict):
                cols = list(data[0].keys())
                return cols, data
        elif isinstance(data, dict):
            # Check if dict of lists (like pandas to_dict('list'))
            first_key = list(data.keys())[0]
            if isinstance(data[first_key], list):
                df = pd.DataFrame(data)
                return list(df.columns), df.to_dict(orient="records")
        raise ValueError("JSON must be an array of objects or dict of lists")
    except Exception as e:
        logger.error("parse_json_failed", error=str(e))
        raise ValueError(f"Failed to parse JSON: {e}")


def parse_docx(contents: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        import docx
        doc = docx.Document(io.BytesIO(contents))
        # Try to find tables first
        if len(doc.tables) > 0:
            table = doc.tables[0]
            rows_data = []
            for row in table.rows:
                text_cells = [cell.text.strip() for cell in row.cells]
                rows_data.append(text_cells)
            
            if len(rows_data) < 2:
                raise ValueError("DOCX table has insufficient data rows")
            
            column_names = rows_data[0]
            column_names = [col if col else f"col_{idx}" for idx, col in enumerate(column_names)]
            
            rows = []
            for r in rows_data[1:]:
                row_dict = {}
                for idx, col in enumerate(column_names):
                    val = r[idx] if idx < len(r) else None
                    row_dict[col] = val
                rows.append(row_dict)
            return column_names, rows
        else:
            # Fallback to paragraph parsing
            lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if not lines:
                raise ValueError("No text or tables found in DOCX file")
            import csv
            reader = csv.DictReader(io.StringIO("\n".join(lines)))
            column_names = reader.fieldnames or []
            rows = [dict(r) for r in reader]
            return column_names, rows
    except Exception as e:
        logger.error("parse_docx_failed", error=str(e))
        raise ValueError(f"Failed to parse DOCX: {e}")


def parse_pdf(contents: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(contents))
        all_text_lines = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                all_text_lines.extend([line.strip() for line in text.split("\n") if line.strip()])
        
        if not all_text_lines:
            raise ValueError("No text extracted from PDF")

        import csv
        sample = "\n".join(all_text_lines[:10])
        if "," in sample:
            csv_reader = csv.DictReader(io.StringIO("\n".join(all_text_lines)))
            column_names = csv_reader.fieldnames or []
            rows = [dict(r) for r in csv_reader]
            if len(column_names) > 1 and len(rows) > 0:
                return column_names, rows

        # Fallback space-separated columns
        first_line = all_text_lines[0]
        column_names = [c.strip() for c in re.split(r'\s{2,}', first_line) if c.strip()]
        if len(column_names) <= 1:
            column_names = [c.strip() for c in first_line.split(" ") if c.strip()]

        rows = []
        for line in all_text_lines[1:]:
            parts = [p.strip() for p in re.split(r'\s{2,}', line) if p.strip()]
            if len(parts) != len(column_names):
                parts = [p.strip() for p in line.split(" ") if p.strip()]
            
            row_dict = {}
            for idx, col in enumerate(column_names):
                val = parts[idx] if idx < len(parts) else None
                row_dict[col] = val
            rows.append(row_dict)
            
        return column_names, rows
    except Exception as e:
        logger.error("parse_pdf_failed", error=str(e))
        raise ValueError(f"Failed to parse PDF: {e}")


def clean_dataset(
    column_names: List[str],
    rows: List[Dict[str, Any]]
) -> Tuple[List[str], List[Dict[str, Any]], Dict[str, Any]]:
    if not rows:
        return column_names, rows, {"cleaned": False, "reason": "No rows to clean"}

    df = pd.DataFrame(rows)
    original_row_count = len(df)

    # 1. Drop duplicate rows
    df = df.drop_duplicates()
    duplicates_removed = original_row_count - len(df)

    # 2. Impute missing values
    imputed_count = 0
    for col in df.columns:
        null_mask = df[col].isna() | (df[col] == "") | (df[col].astype(str).str.lower() == "null")
        null_count = null_mask.sum()
        if null_count > 0:
            imputed_count += int(null_count)
            num_series = pd.to_numeric(df[col], errors="coerce")
            if pd.api.types.is_numeric_dtype(num_series.dropna()) and num_series.notna().sum() > 0:
                median_val = num_series.median()
                fill_val = median_val if pd.notna(median_val) else 0
                df.loc[null_mask, col] = fill_val
            else:
                df.loc[null_mask, col] = "Unknown"

    cleaned_rows = df.to_dict(orient="records")
    summary = {
        "original_rows": original_row_count,
        "cleaned_rows": len(cleaned_rows),
        "duplicates_removed": duplicates_removed,
        "imputed_missing_values": imputed_count,
        "is_cleaned": True
    }
    return list(df.columns), cleaned_rows, summary
