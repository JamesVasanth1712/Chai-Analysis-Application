from typing import Any, Dict, List, Optional
from fastapi import APIRouter, Depends, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete, update
from pydantic import BaseModel
from app.api.deps import get_current_user, get_current_tenant
from app.core.database import get_db
from app.models.user import User
from app.models.tenant import Tenant
from app.models.dashboard import Dashboard, DashboardWidget
from app.models.extensions import DashboardTemplate
import structlog

logger = structlog.get_logger()
router = APIRouter(prefix="/dashboards", tags=["dashboards"])


class WidgetCreate(BaseModel):
    title: str
    chart_type: str = "bar"
    query_sql: Optional[str] = None
    query_nl: Optional[str] = None
    x_key: Optional[str] = None
    y_keys: Optional[List[str]] = None
    position_x: int = 0
    position_y: int = 0
    width: int = 6
    height: int = 4
    refresh_interval_s: Optional[int] = None
    date_filter_enabled: bool = False
    dataset_id: Optional[str] = None


class DashboardCreate(BaseModel):
    title: str
    description: Optional[str] = None
    is_default: bool = False
    widgets: Optional[List[WidgetCreate]] = None


class DashboardUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    is_default: Optional[bool] = None
    layout_config: Optional[Dict[str, Any]] = None


class LayoutUpdateItem(BaseModel):
    id: str
    position_x: int
    position_y: int
    width: int
    height: int


@router.post("/")
async def create_dashboard(
    body: DashboardCreate,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    if body.is_default:
        result = await db.execute(
            select(Dashboard).where(
                Dashboard.tenant_id == current_tenant.id,
                Dashboard.is_default == True,
            )
        )
        existing = result.scalars().all()
        for d in existing:
            d.is_default = False

    dashboard = Dashboard(
        title=body.title,
        description=body.description,
        is_default=body.is_default,
        layout_config={},
        tenant_id=current_tenant.id,
        created_by=current_user.id,
    )
    db.add(dashboard)
    await db.flush()

    for w in (body.widgets or []):
        widget = DashboardWidget(
            dashboard_id=dashboard.id,
            title=w.title,
            chart_type=w.chart_type,
            query_sql=w.query_sql,
            query_nl=w.query_nl,
            x_key=w.x_key,
            y_keys=w.y_keys,
            position_x=w.position_x,
            position_y=w.position_y,
            width=w.width,
            height=w.height,
            refresh_interval_s=w.refresh_interval_s,
            date_filter_enabled=w.date_filter_enabled,
            dataset_id=w.dataset_id,
        )
        db.add(widget)

    await db.commit()
    return {"id": dashboard.id, "title": dashboard.title}


@router.get("/")
async def list_dashboards(
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard)
        .where(Dashboard.tenant_id == current_tenant.id)
        .order_by(Dashboard.is_default.desc(), Dashboard.created_at.desc())
    )
    dashboards = result.scalars().all()
    return [
        {
            "id": d.id,
            "title": d.title,
            "description": d.description,
            "is_default": d.is_default,
            "created_at": d.created_at.isoformat(),
        }
        for d in dashboards
    ]


@router.get("/{dashboard_id}")
async def get_dashboard(
    dashboard_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    dashboard = result.scalar_one_or_none()
    if not dashboard:
        raise HTTPException(status_code=404, detail="Dashboard not found")

    widget_result = await db.execute(
        select(DashboardWidget)
        .where(DashboardWidget.dashboard_id == dashboard_id)
        .order_by(DashboardWidget.position_y, DashboardWidget.position_x)
    )
    widgets = widget_result.scalars().all()

    return {
        "id": dashboard.id,
        "title": dashboard.title,
        "description": dashboard.description,
        "is_default": dashboard.is_default,
        "layout_config": dashboard.layout_config,
        "widgets": [
            {
                "id": w.id,
                "title": w.title,
                "chart_type": w.chart_type,
                "query_sql": w.query_sql,
                "query_nl": w.query_nl,
                "x_key": w.x_key,
                "y_keys": w.y_keys,
                "position_x": w.position_x,
                "position_y": w.position_y,
                "width": w.width,
                "height": w.height,
                "refresh_interval_s": w.refresh_interval_s,
                "date_filter_enabled": w.date_filter_enabled,
                "dataset_id": w.dataset_id,
            }
            for w in widgets
        ],
        "created_at": dashboard.created_at.isoformat(),
    }


@router.patch("/{dashboard_id}")
async def update_dashboard(
    dashboard_id: str,
    body: DashboardUpdate,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    dashboard = result.scalar_one_or_none()
    if not dashboard:
        raise HTTPException(status_code=404, detail="Dashboard not found")

    if body.is_default:
        others = await db.execute(
            select(Dashboard).where(
                Dashboard.tenant_id == current_tenant.id,
                Dashboard.id != dashboard_id,
                Dashboard.is_default == True,
            )
        )
        for d in others.scalars().all():
            d.is_default = False

    if body.title is not None:
        dashboard.title = body.title
    if body.description is not None:
        dashboard.description = body.description
    if body.is_default is not None:
        dashboard.is_default = body.is_default
    if body.layout_config is not None:
        dashboard.layout_config = body.layout_config

    await db.commit()
    return {"id": dashboard_id, "ok": True}


@router.delete("/{dashboard_id}")
async def delete_dashboard(
    dashboard_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    dashboard = result.scalar_one_or_none()
    if not dashboard:
        raise HTTPException(status_code=404, detail="Dashboard not found")

    from app.models.dataset import Dataset

    await db.execute(
        update(Dataset)
        .where(Dataset.dashboard_id == dashboard_id, Dataset.tenant_id == current_tenant.id)
        .values(dashboard_id=None)
    )
    await db.execute(delete(DashboardWidget).where(DashboardWidget.dashboard_id == dashboard_id))
    await db.execute(delete(Dashboard).where(Dashboard.id == dashboard_id))
    await db.commit()
    return {"ok": True, "id": dashboard_id}


# Widget endpoints
@router.post("/{dashboard_id}/widgets")
async def add_widget(
    dashboard_id: str,
    body: WidgetCreate,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    if not result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Dashboard not found")

    widget = DashboardWidget(
        dashboard_id=dashboard_id,
        title=body.title,
        chart_type=body.chart_type,
        query_sql=body.query_sql,
        query_nl=body.query_nl,
        x_key=body.x_key,
        y_keys=body.y_keys,
        position_x=body.position_x,
        position_y=body.position_y,
        width=body.width,
        height=body.height,
        refresh_interval_s=body.refresh_interval_s,
        date_filter_enabled=body.date_filter_enabled,
        dataset_id=body.dataset_id,
    )
    db.add(widget)
    await db.commit()
    return {"id": widget.id, "dashboard_id": dashboard_id}


@router.patch("/{dashboard_id}/widgets/{widget_id}")
async def update_widget(
    dashboard_id: str,
    widget_id: str,
    body: WidgetCreate,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(DashboardWidget).where(
            DashboardWidget.id == widget_id,
            DashboardWidget.dashboard_id == dashboard_id,
        )
    )
    widget = result.scalar_one_or_none()
    if not widget:
        raise HTTPException(status_code=404, detail="Widget not found")

    widget.title = body.title
    widget.chart_type = body.chart_type
    widget.query_sql = body.query_sql
    widget.query_nl = body.query_nl
    widget.x_key = body.x_key
    widget.y_keys = body.y_keys
    widget.position_x = body.position_x
    widget.position_y = body.position_y
    widget.width = body.width
    widget.height = body.height
    widget.refresh_interval_s = body.refresh_interval_s
    widget.date_filter_enabled = body.date_filter_enabled
    widget.dataset_id = body.dataset_id
    await db.commit()
    return {"id": widget_id, "ok": True}


@router.delete("/{dashboard_id}/widgets/{widget_id}")
async def delete_widget(
    dashboard_id: str,
    widget_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(DashboardWidget).where(
            DashboardWidget.id == widget_id,
            DashboardWidget.dashboard_id == dashboard_id,
        )
    )
    if not result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Widget not found")

    await db.execute(
        delete(DashboardWidget).where(DashboardWidget.id == widget_id)
    )
    await db.commit()
    return {"ok": True, "id": widget_id}


@router.put("/{dashboard_id}/layout")
async def update_layout(
    dashboard_id: str,
    items: List[LayoutUpdateItem],
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    if not result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Dashboard not found")

    for item in items:
        widget_result = await db.execute(
            select(DashboardWidget).where(
                DashboardWidget.id == item.id,
                DashboardWidget.dashboard_id == dashboard_id,
            )
        )
        widget = widget_result.scalar_one_or_none()
        if widget:
            widget.position_x = item.position_x
            widget.position_y = item.position_y
            widget.width = item.width
            widget.height = item.height

    await db.commit()
    return {"ok": True}


@router.get("/{dashboard_id}/widgets/{widget_id}/data")
async def get_widget_data(
    dashboard_id: str,
    widget_id: str,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(DashboardWidget).where(
            DashboardWidget.id == widget_id,
            DashboardWidget.dashboard_id == dashboard_id,
        )
    )
    widget = result.scalar_one_or_none()
    if not widget:
        raise HTTPException(status_code=404, detail="Widget not found")

    from app.services.nl_to_sql import run_query, shape_for_chart
    db_url = (current_tenant.settings or {}).get("external_db_url")

    dataset_table = None
    if widget.dataset_id:
        from app.models.dataset import Dataset as DatasetModel
        ds_result = await db.execute(
            select(DatasetModel).where(
                DatasetModel.id == widget.dataset_id,
                DatasetModel.tenant_id == current_tenant.id,
            )
        )
        ds = ds_result.scalar_one_or_none()
        if ds:
            dataset_table = ds.table_name

    sql = widget.query_sql
    if dataset_table and sql:
        sql = sql.replace("dataset", f'"{dataset_table}"')
    elif dataset_table and not sql and not widget.query_nl:
        sql = f'SELECT * FROM "{dataset_table}"'

    # Load dashboard layout_config to fetch calculated columns & relationships
    dashboard_res = await db.execute(
        select(Dashboard).where(
            Dashboard.id == dashboard_id,
            Dashboard.tenant_id == current_tenant.id,
        )
    )
    dashboard = dashboard_res.scalar_one_or_none()
    calculated_columns = []
    relationships = []
    if dashboard and dashboard.layout_config:
        calculated_columns = dashboard.layout_config.get("calculated_columns", [])
        relationships = dashboard.layout_config.get("relationships", [])

    # Fetch all datasets
    datasets_res = await db.execute(
        select(DatasetModel).where(DatasetModel.tenant_id == current_tenant.id)
    )
    all_datasets = list(datasets_res.scalars().all())

    # Rewrite query SQL for modeling (calculated columns & relationships join)
    from app.services.analysis_service import rewrite_sql_for_modeling
    if sql:
        sql = rewrite_sql_for_modeling(sql, dataset_table, calculated_columns, relationships, all_datasets)

    # Sanitize dataset names and fix corrupted aggregate aliases before execution
    from app.api.routes.data import rewrite_dataset_names_in_sql, sanitize_aliased_sql
    if sql:
        sql = rewrite_dataset_names_in_sql(sql, all_datasets)
        sql = sanitize_aliased_sql(sql)

    try:
        query_result = await run_query(
            sql=sql,
            nl_query=widget.query_nl if not sql else None,
            tenant_id=current_tenant.id,
            db_url=db_url,
            limit=500,
        )
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    chart_data = shape_for_chart(
        query_result["rows"], query_result["columns"],
        widget.x_key, widget.y_keys, widget.chart_type
    )
    return {
        "widget_id": widget_id,
        "chart_type": widget.chart_type,
        "sql": query_result["sql"],
        "row_count": query_result["row_count"],
        **chart_data,
    }


class ConnectionManager:
    def __init__(self):
        # Maps dashboard_id -> list of (websocket, user_id, user_name)
        self.active_connections: Dict[str, List[tuple]] = {}

    async def connect(self, dashboard_id: str, websocket: WebSocket, user_id: str, user_name: str):
        await websocket.accept()
        if dashboard_id not in self.active_connections:
            self.active_connections[dashboard_id] = []
        self.active_connections[dashboard_id].append((websocket, user_id, user_name))

    def disconnect(self, dashboard_id: str, websocket: WebSocket):
        if dashboard_id in self.active_connections:
            self.active_connections[dashboard_id] = [
                conn for conn in self.active_connections[dashboard_id] if conn[0] != websocket
            ]
            if not self.active_connections[dashboard_id]:
                del self.active_connections[dashboard_id]

    async def broadcast(self, dashboard_id: str, message: dict, exclude_websocket: WebSocket = None):
        if dashboard_id in self.active_connections:
            for conn, user_id, user_name in self.active_connections[dashboard_id]:
                if conn != exclude_websocket:
                    try:
                        await conn.send_json(message)
                    except Exception:
                        pass

manager = ConnectionManager()


@router.websocket("/ws/{dashboard_id}")
async def websocket_endpoint(
    websocket: WebSocket,
    dashboard_id: str,
    user_id: str,
    user_name: str,
):
    await manager.connect(dashboard_id, websocket, user_id, user_name)
    # Broadcast that user joined
    await manager.broadcast(
        dashboard_id,
        {"type": "user_joined", "user_id": user_id, "user_name": user_name},
        exclude_websocket=websocket
    )
    try:
        while True:
            data = await websocket.receive_json()
            # Broadcast cursor movement, card lock, or card unlock
            await manager.broadcast(
                dashboard_id,
                {"user_id": user_id, "user_name": user_name, **data},
                exclude_websocket=websocket
            )
    except WebSocketDisconnect:
        manager.disconnect(dashboard_id, websocket)
        await manager.broadcast(
            dashboard_id,
            {"type": "user_left", "user_id": user_id, "user_name": user_name}
        )


@router.get("/{dashboard_id}/export/excel")
async def export_excel(
    dashboard_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    import pandas as pd
    import io
    from app.models.dataset import Dataset as DatasetModel
    
    widget_result = await db.execute(
        select(DashboardWidget).where(DashboardWidget.dashboard_id == dashboard_id)
    )
    widgets = widget_result.scalars().all()
    
    dataset_ids = {w.dataset_id for w in widgets if w.dataset_id}
    if not dataset_ids:
        raise HTTPException(status_code=400, detail="No datasets found for this dashboard")
        
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for ds_id in dataset_ids:
            ds_res = await db.execute(
                select(DatasetModel).where(DatasetModel.id == ds_id, DatasetModel.tenant_id == current_tenant.id)
            )
            ds = ds_res.scalar_one_or_none()
            if ds and ds.table_name:
                from app.tools.sql_executor import sql_executor
                rows = await sql_executor.execute(f'SELECT * FROM "{ds.table_name}"', limit=10000)
                if rows:
                    df = pd.DataFrame(rows)
                    sheet_name = ds.name[:30]
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=dashboard_data_export.xlsx"}
    )


@router.get("/{dashboard_id}/export/docx")
async def export_docx(
    dashboard_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    import io
    import docx
    from app.models.dataset import Dataset as DatasetModel
    
    dashboard_res = await db.execute(
        select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.tenant_id == current_tenant.id)
    )
    dashboard = dashboard_res.scalar_one_or_none()
    if not dashboard:
        raise HTTPException(status_code=404, detail="Dashboard not found")

    widget_result = await db.execute(
        select(DashboardWidget).where(DashboardWidget.dashboard_id == dashboard_id)
    )
    widgets = widget_result.scalars().all()

    doc = docx.Document()
    doc.add_heading(dashboard.title, 0)
    doc.add_paragraph("Report workspace export generated by Chai Analysis")
    doc.add_paragraph(f"Created by: {current_user.full_name}")
    doc.add_paragraph(f"Date: {dashboard.created_at.strftime('%Y-%m-%d %H:%M:%S')}")
    
    doc.add_heading("Widgets and Visualizations", level=1)
    for widget in widgets:
        doc.add_heading(widget.title or "Visual Card", level=2)
        doc.add_paragraph(f"Chart Type: {widget.chart_type}")
        doc.add_paragraph("Query SQL:")
        doc.add_paragraph(widget.query_sql or "No SQL", style="Normal")
        doc.add_paragraph("")
        
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={dashboard.title.replace(' ', '_')}_report.docx"}
    )


@router.get("/{dashboard_id}/export/pptx")
async def export_pptx(
    dashboard_id: str,
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
    db: AsyncSession = Depends(get_db),
):
    import io
    import pptx
    from app.models.dataset import Dataset as DatasetModel
    
    dashboard_res = await db.execute(
        select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.tenant_id == current_tenant.id)
    )
    dashboard = dashboard_res.scalar_one_or_none()
    if not dashboard:
        raise HTTPException(status_code=404, detail="Dashboard not found")

    widget_result = await db.execute(
        select(DashboardWidget).where(DashboardWidget.dashboard_id == dashboard_id)
    )
    widgets = widget_result.scalars().all()

    prs = pptx.Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = dashboard.title
    subtitle.text = f"Report workspace export\nCreated by {current_user.full_name}\nDate: {dashboard.created_at.strftime('%Y-%m-%d')}"

    # slide for each widget
    blank_slide_layout = prs.slide_layouts[6]
    for widget in widgets:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Add Title Box
        txBox = slide.shapes.add_textbox(pptx.util.Inches(0.5), pptx.util.Inches(0.5), pptx.util.Inches(9), pptx.util.Inches(1))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f"Visual: {widget.title}"
        p.font.size = pptx.util.Pt(28)
        p.font.bold = True
        
        # Add Description Box
        descBox = slide.shapes.add_textbox(pptx.util.Inches(0.5), pptx.util.Inches(1.5), pptx.util.Inches(9), pptx.util.Inches(1))
        tf_desc = descBox.text_frame
        p_desc = tf_desc.paragraphs[0]
        p_desc.text = f"Chart Type: {widget.chart_type.upper()}  |  Date Filter: {'Enabled' if widget.date_filter_enabled else 'Disabled'}"
        p_desc.font.size = pptx.util.Pt(14)
        p_desc.font.italic = True
        
        # Add SQL Box
        sqlBox = slide.shapes.add_textbox(pptx.util.Inches(0.5), pptx.util.Inches(2.5), pptx.util.Inches(9), pptx.util.Inches(3))
        tf_sql = sqlBox.text_frame
        tf_sql.word_wrap = True
        p_sql = tf_sql.paragraphs[0]
        p_sql.text = f"Query SQL:\n{widget.query_sql or 'No SQL'}"
        p_sql.font.size = pptx.util.Pt(11)
        p_sql.font.name = "Courier New"

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename={dashboard.title.replace(' ', '_')}_export.pptx"}
    )


@router.get("/templates/all")
async def list_templates(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(select(DashboardTemplate))
    templates = result.scalars().all()
    return [
        {
            "id": t.id,
            "name": t.name,
            "category": t.category,
            "description": t.description,
            "layout_config": t.layout_config,
        }
        for t in templates
    ]


class CreateFromTemplateRequest(BaseModel):
    template_id: str
    dataset_id: str
    title: Optional[str] = None


@router.post("/create-from-template")
async def create_from_template(
    body: CreateFromTemplateRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
):
    import pandas as pd
    from app.models.dataset import Dataset
    from app.tools.sql_executor import sql_executor
    
    template_res = await db.execute(
        select(DashboardTemplate).where(DashboardTemplate.id == body.template_id)
    )
    template = template_res.scalar_one_or_none()
    if not template:
        raise HTTPException(status_code=404, detail="Template not found")

    dataset_res = await db.execute(
        select(Dataset).where(Dataset.id == body.dataset_id, Dataset.tenant_id == current_tenant.id)
    )
    dataset = dataset_res.scalar_one_or_none()
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")

    rows = await sql_executor.execute(f'SELECT * FROM "{dataset.table_name}" LIMIT 100')
    if not rows:
        raise HTTPException(status_code=400, detail="The dataset has no records to profile.")
    df = pd.DataFrame(rows)

    # Profile columns
    numeric = [col for col in df.columns if pd.api.types.is_numeric_dtype(df[col])]
    metrics = [col for col in numeric if not any(x in col.lower() for x in ["id", "s_no", "serial", "postcode", "zip"])]
    if not metrics:
        metrics = numeric
        
    categorical = [col for col in df.columns if not pd.api.types.is_numeric_dtype(df[col]) and df[col].nunique(dropna=True) <= 50]
    cats = [col for col in categorical if not any(x in col.lower() for x in ["id", "postcode"])] + [col for col in categorical if any(x in col.lower() for x in ["id", "postcode"])]
    
    dates = []
    for col in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[col]) or any(x in col.lower() for x in ["date", "time", "year", "month"]):
            dates.append(col)

    dashboard = Dashboard(
        title=body.title or f"{template.name} ({dataset.name})",
        description=template.description,
        is_default=False,
        layout_config={"generated_from_template_id": template.id, "pages": ["Page 1"]},
        tenant_id=current_tenant.id,
        created_by=current_user.id,
    )
    db.add(dashboard)
    await db.flush()

    widgets_to_add = []
    x_offset = 0
    y_offset = 0
    
    for idx, widget_spec in enumerate(template.layout_config.get("widgets", [])):
        chart_type = widget_spec.get("chart_type", "bar")
        title = widget_spec.get("title", "Chart")
        slot = widget_spec.get("slot")
        metric_slot = widget_spec.get("metric_slot")
        agg_func = widget_spec.get("agg_func", "SUM")
        
        # Resolve metric
        met_col = None
        if metric_slot == "primary_metric" or slot == "primary_metric":
            met_col = metrics[0] if metrics else None
        elif metric_slot == "secondary_metric" or slot == "secondary_metric":
            met_col = metrics[1] if len(metrics) > 1 else (metrics[0] if metrics else None)

        # Resolve category
        cat_col = None
        if slot == "category_1":
            cat_col = cats[0] if cats else None
        elif slot == "category_2":
            cat_col = cats[1] if len(cats) > 1 else (cats[0] if cats else None)

        # Resolve date
        date_col = None
        if slot == "date_1":
            date_col = dates[0] if dates else None

        sql_query = None
        x_key = None
        y_keys = ["y"]

        if chart_type == "card":
            if slot == "row_count":
                sql_query = f'SELECT COUNT(*) AS "y" FROM "{dataset.table_name}"'
            elif met_col:
                sql_query = f'SELECT {agg_func}("{met_col}") AS "y" FROM "{dataset.table_name}"'
            elif cat_col:
                sql_query = f'SELECT COUNT(DISTINCT "{cat_col}") AS "y" FROM "{dataset.table_name}"'
            else:
                sql_query = f'SELECT COUNT(*) AS "y" FROM "{dataset.table_name}"'
            x_key = None
            
        elif chart_type in ["bar", "pie", "area"]:
            active_cat = cat_col or (cats[0] if cats else None)
            active_met = met_col or (metrics[0] if metrics else None)
            if active_cat:
                if active_met:
                    sql_query = f'SELECT "{active_cat}" AS "x", {agg_func}("{active_met}") AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 2 DESC LIMIT 10'
                else:
                    sql_query = f'SELECT "{active_cat}" AS "x", COUNT(*) AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 2 DESC LIMIT 10'
                x_key = "x"
            else:
                sql_query = f'SELECT \'No Category\' AS "x", 0 AS "y"'
                x_key = "x"
                
        elif chart_type == "line":
            active_date = date_col or (dates[0] if dates else None)
            active_met = met_col or (metrics[0] if metrics else None)
            if active_date:
                is_datetime = df[active_date].dtype == "datetime64[ns]" or "date" in active_date.lower()
                if active_met:
                    if is_datetime:
                        sql_query = f'SELECT DATE_TRUNC(\'day\', "{active_date}") AS "x", {agg_func}("{active_met}") AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 1 ASC LIMIT 100'
                    else:
                        sql_query = f'SELECT "{active_date}" AS "x", {agg_func}("{active_met}") AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 1 ASC LIMIT 100'
                else:
                    if is_datetime:
                        sql_query = f'SELECT DATE_TRUNC(\'day\', "{active_date}") AS "x", COUNT(*) AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 1 ASC LIMIT 100'
                    else:
                        sql_query = f'SELECT "{active_date}" AS "x", COUNT(*) AS "y" FROM "{dataset.table_name}" GROUP BY 1 ORDER BY 1 ASC LIMIT 100'
                x_key = "x"
            else:
                sql_query = f'SELECT \'No Timeline\' AS "x", 0 AS "y"'
                x_key = "x"

        w_width = widget_spec.get("width", 6)
        w_height = widget_spec.get("height", 4)
        
        widget = DashboardWidget(
            dashboard_id=dashboard.id,
            title=title,
            chart_type=chart_type,
            query_sql=sql_query,
            x_key=x_key,
            y_keys=y_keys,
            position_x=x_offset,
            position_y=y_offset,
            width=w_width,
            height=w_height,
            dataset_id=dataset.id,
            date_filter_enabled=True
        )
        db.add(widget)
        
        x_offset += w_width
        if x_offset >= 12:
            x_offset = 0
            y_offset += w_height

    await db.commit()
    return {"status": "success", "dashboard_id": dashboard.id}


class ScheduleCreate(BaseModel):
    frequency: str
    recipients: List[str]
    is_active: bool = True


@router.post("/{dashboard_id}/schedule")
async def create_schedule(
    dashboard_id: str,
    body: ScheduleCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
):
    from app.models.extensions import ScheduledReport
    # Check if dashboard exists
    dashboard_res = await db.execute(
        select(Dashboard).where(Dashboard.id == dashboard_id, Dashboard.tenant_id == current_tenant.id)
    )
    if not dashboard_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Dashboard not found")

    report = ScheduledReport(
        frequency=body.frequency,
        recipients=body.recipients,
        is_active=body.is_active,
        dashboard_id=dashboard_id,
        tenant_id=current_tenant.id,
    )
    db.add(report)
    await db.commit()
    return {"status": "success", "schedule_id": report.id}


@router.get("/{dashboard_id}/schedules")
async def get_schedules(
    dashboard_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
):
    from app.models.extensions import ScheduledReport
    result = await db.execute(
        select(ScheduledReport).where(
            ScheduledReport.dashboard_id == dashboard_id,
            ScheduledReport.tenant_id == current_tenant.id
        )
    )
    schedules = result.scalars().all()
    return [
        {
            "id": s.id,
            "frequency": s.frequency,
            "recipients": s.recipients,
            "is_active": s.is_active,
            "created_at": s.created_at.isoformat() if s.created_at else None,
            "last_run_at": s.last_run_at.isoformat() if s.last_run_at else None,
        }
        for s in schedules
    ]


@router.delete("/schedules/{schedule_id}")
async def delete_schedule(
    schedule_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
    current_tenant: Tenant = Depends(get_current_tenant),
):
    from app.models.extensions import ScheduledReport
    result = await db.execute(
        select(ScheduledReport).where(
            ScheduledReport.id == schedule_id,
            ScheduledReport.tenant_id == current_tenant.id
        )
    )
    schedule = result.scalar_one_or_none()
    if not schedule:
        raise HTTPException(status_code=404, detail="Schedule not found")

    await db.delete(schedule)
    await db.commit()
    return {"status": "success"}
